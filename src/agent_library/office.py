"""Local Office worker: MarkItDown rendering with retained tables and embedded images.

Only DOCX/XLSX are enabled. No models, generic URL handling or plugins are used.
"""
import html
import io
import json
import sys
import zipfile
from importlib.metadata import version
from pathlib import Path

if __package__:
    from .artifacts import add_asset, table_html
else:
    from artifacts import add_asset, table_html


def image_asset(result, data, media_type):
    formats = {"image/png": ("png", b"\x89PNG\r\n\x1a\n"), "image/jpeg": ("jpg", b"\xff\xd8"),
               "image/gif": ("gif", b"GIF"), "image/bmp": ("bmp", b"BM")}
    extension, magic = formats.get(media_type, ("bin", b""))
    if extension == "bin" or not data.startswith(magic):
        result["warnings"].append("IMAGE_PRESERVED_AS_ATTACHMENT")
        result["status"] = "partial"
        extension, media_type = "bin", "application/octet-stream"
    return add_asset(result, data, extension, media_type, "embedded_image")


def html_page(result, markup, number, locator):
    from bs4 import BeautifulSoup
    from markitdown import StreamInfo
    from markitdown.converters import HtmlConverter
    soup = BeautifulSoup(markup, "html.parser")
    page = {"number": number, "locator": locator, "text": soup.get_text("\n"), "images": [], "tables": []}
    replacements = {}
    for index, table in enumerate(list(soup.find_all("table")), 1):
        if table.find_parent("table") is not None:
            continue
        if table.find("table"):
            result["status"] = "partial"
            result["warnings"].append("NESTED_TABLE_REQUIRES_REVIEW")
        cells, occupied, cols = [], set(), 0
        rows = [r for r in table.find_all("tr") if r.find_parent("table") is table]
        for r, row in enumerate(rows):
            c = 0
            for td in row.find_all(["td", "th"], recursive=False):
                while (r, c) in occupied:
                    c += 1
                rs, cs = int(td.get("rowspan", 1)), int(td.get("colspan", 1))
                if not (1 <= rs <= 10000 and 1 <= cs <= 1000):
                    raise ValueError("Invalid table span")
                cells.append({"row": r, "col": c, "rowspan": rs, "colspan": cs,
                              "text": td.get_text("\n"), "header": td.name == "th"})
                occupied.update((a, b) for a in range(r, r + rs) for b in range(c, c + cs))
                c += cs
                cols = max(cols, c)
        rendered = table_html(cells, len(rows), cols)
        asset = add_asset(result, rendered.encode(), "html", "text/html", "structured_table")
        page["tables"].append({"id": f"table-{number}-{index}", "locator": {**locator, "table_index": index},
                               "rows": len(rows), "cols": cols, "cells": cells, "asset": asset})
        token = f"AGENTLIBRARYTABLE{number}X{index}PLACEHOLDER"
        while token in markup:
            token += "X"
        replacements[token] = rendered
        # Keep image references even if an image was inside a table cell.
        for img in table.find_all("img"):
            if img.get("src") in result["assets"]:
                page["images"].append(img["src"])
                replacements[token] += f"\n\n![Table image]({img['src']})\n"
        table.replace_with(soup.new_string(token))
    for tag in list(soup.find_all(True)):
        if tag.name == "img":
            src = tag.get("src")
            if src in result["assets"]:
                page["images"].append(src)
                tag.attrs = {"src": src, "alt": tag.get("alt") or "Embedded source image"}
            else:
                result["status"] = "partial"
                result["warnings"].append("EXTERNAL_OR_MISSING_IMAGE")
                tag.replace_with("[External or missing image; inspect original]")
        elif tag.name not in {"p", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "li", "br", "strong", "em", "code", "pre", "blockquote"}:
            tag.unwrap()
        else:
            tag.attrs = {}
    converted = HtmlConverter().convert(io.BytesIO(str(soup).encode()), StreamInfo(extension=".html", charset="utf-8"))
    page["markdown"] = converted.markdown
    for token, rendered in replacements.items():
        page["markdown"] = page["markdown"].replace(token, rendered)
    return page


def convert(source):
    if version("markitdown") != "0.1.7":
        raise ValueError("Unverified MarkItDown version")
    result = {"schema_version": 2, "method": "markitdown-0.1.7:office-v1", "status": "extracted",
              "pages": [], "assets": {}, "warnings": [],
              "capabilities": {"visuals": "embedded_images", "tables": "native_cells_and_spans", "ocr": False,
                               "page_basis": "logical_sections_not_printed_pages"},
              "dependencies": {n: version(n) for n in ["markitdown", "mammoth", "openpyxl", "markdownify"]}}
    with zipfile.ZipFile(source) as archive:
        infos = archive.infolist()
        if len(infos) > 4096 or sum(i.file_size for i in infos) > 100 * 1024 * 1024:
            raise ValueError("Office container exceeds limits")
        if len({i.filename for i in infos}) != len(infos):
            raise ValueError("Duplicate Office parts")
        # Preserve every packaged image, including those in headers or unsupported drawings.
        media = []
        for info in infos:
            if info.filename.startswith(("word/media/", "xl/media/")) and not info.is_dir():
                import mimetypes
                mime = mimetypes.guess_type(info.filename)[0] or "application/octet-stream"
                name = image_asset(result, archive.read(info), mime)
                media.append({"asset": name, "package_part": info.filename})
        result["embedded_media"] = media
        if source.suffix == ".docx":
            import mammoth
            if any(i.filename.startswith(("word/header", "word/footer", "word/charts/", "word/diagrams/",
                                           "word/embeddings/", "word/comments")) for i in infos):
                result["warnings"].append("OFFICE_PARTS_REQUIRE_ORIGINAL")
                result["status"] = "partial"
            result["raw_asset"] = add_asset(result, archive.read("word/document.xml"), "xml", "application/xml", "native_document_structure")
            def save_image(image):
                with image.open() as stream:
                    return {"src": image_asset(result, stream.read(), image.content_type)}
            with source.open("rb") as stream:
                converted = mammoth.convert_to_html(stream, convert_image=mammoth.images.img_element(save_image),
                                                    external_file_access=False, include_embedded_style_map=False)
            if converted.messages:
                result["warnings"].append("OFFICE_CONVERSION_WARNINGS")
                result["status"] = "partial"
            result["pages"] = [html_page(result, converted.value, 1, {"package_part": "word/document.xml"})]
        else:
            from openpyxl import load_workbook
            # Never evaluate formulas or follow external workbook links.
            workbook = load_workbook(source, data_only=False, keep_links=False)
            cached = load_workbook(source, data_only=True, keep_links=False)
            native = []
            try:
                for number, sheet in enumerate(workbook.worksheets, 1):
                    if sheet.max_row * sheet.max_column > 100000:
                        raise ValueError("Worksheet exceeds limits")
                    spans = {(m.min_row, m.min_col): m for m in sheet.merged_cells.ranges}
                    covered = {(r, c) for m in sheet.merged_cells.ranges for r in range(m.min_row, m.max_row + 1)
                               for c in range(m.min_col, m.max_col + 1) if (r, c) != (m.min_row, m.min_col)}
                    cells, values = [], []
                    for row in sheet.iter_rows():
                        for cell in row:
                            if (cell.row, cell.column) in covered:
                                continue
                            span = spans.get((cell.row, cell.column))
                            value = "" if cell.value is None else str(cell.value)
                            cells.append({"row": cell.row - 1, "col": cell.column - 1, "text": value,
                                "rowspan": span.max_row - span.min_row + 1 if span else 1,
                                "colspan": span.max_col - span.min_col + 1 if span else 1})
                            if cell.value is not None:
                                cached_value = cached[sheet.title][cell.coordinate].value
                                values.append({"address": cell.coordinate, "value": value, "data_type": cell.data_type,
                                    "number_format": cell.number_format, "cached_value": None if cached_value is None else str(cached_value)})
                    markup = "<h1>" + html.escape(sheet.title) + "</h1>" + table_html(cells, sheet.max_row, sheet.max_column)
                    result["pages"].append(html_page(result, markup, number, {"sheet": sheet.title}))
                    native.append({"sheet": sheet.title, "state": sheet.sheet_state, "cells": values,
                                   "merged_ranges": [str(m) for m in sheet.merged_cells.ranges]})
                result["raw_asset"] = add_asset(result, json.dumps(native, ensure_ascii=False, sort_keys=True).encode(),
                                                "json", "application/json", "native_workbook_cells")
                result["warnings"].append("WORKBOOK_CHARTS_AND_DRAWINGS_REQUIRE_ORIGINAL")
                result["status"] = "partial"  # No renderer for charts, shapes or print layout in this profile.
            finally:
                workbook.close()
                cached.close()
        located = {name for page in result["pages"] for name in page["images"]}
        if any(m["asset"] not in located for m in media):
            result["warnings"].append("UNLOCATED_EMBEDDED_MEDIA")
            result["status"] = "partial"
    if not any(p["text"].strip() for p in result["pages"]):
        result["status"] = "needs_ocr"
        result["warnings"].append("NO_EXTRACTED_TEXT")
    result["warnings"] = sorted(set(result["warnings"]))
    return result


if __name__ == "__main__":
    source, output = map(Path, sys.argv[1:])
    result = convert(source)
    for name, data in result.pop("_asset_bytes", {}).items():
        path = output / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
    (output / "result.json").write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
